"""CPM Import Service (§2.3–2.4, §3.21).

The CPM ``.xlsx`` is the single master source for infrastructure data. This
service:

  * auto-detects the header row and maps columns by **name** (Persian *or*
    standardized English), so column re-ordering in the monthly file is safe;
  * normalises packed / localized values (``2G3G4G`` → flags, ``هدف`` → Target,
    ``No Site ID`` → temporary-code fallback);
  * groups rows into Site → Work Item → Village;
  * applies the sync rules: **new** → create, **unchanged** → ignore,
    **changed** → raise a Change Request for PM decision (records are never
    silently overwritten).

No Work Item is ever created outside this path (invariant §3.22.11 / §2.3).
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

import openpyxl
from sqlalchemy.orm import Session

from app.models.audit import ChangeRequest, CpmImportBatch
from app.models.enums import (
    ApprovalStatus,
    SiteType,
    Technology,
    VillageType,
)
from app.models.master import City, Province
from app.models.site import NO_SITE_ID, Site
from app.models.village import Village, VillageTechApproval
from app.models.work_item import WorkItem

# --- Column aliases: canonical field -> accepted header names (normalized) ---
COLUMN_ALIASES: dict[str, list[str]] = {
    "village_id": ["village_id", "village id", "کد آبادی"],
    "province": ["province", "استان"],
    "city": ["city", "شهرستان"],
    "village_name": ["village_name", "village name", "آبادی"],
    "site_id": ["site_id", "site id", "کد سایت ایرانسل"],
    "temp_site_id": ["temp_site_id", "temp site id", "کدسایت موقت", "کد سایت موقت"],
    "site_assignment_date": [
        "site_assignment_date", "site assignment date",
        "تاریخ ارجاع (میلادی)", "تاریخ ارجاع (شمسی)",
    ],
    "site_type": ["site_type", "site type", "نوع سایت"],
    "requested_technology": [
        "requested_technology", "requested technology", "تکنولوژی درخواستی",
    ],
    "on_air_date": [
        "on_air_date", "on-air date", "on air date",
        "تاریخ راه اندازی (شمسی)", "تاریخ راه اندازی (میلادی)",
    ],
    "last_site_status": [
        "last_site_status", "last site status", "آخرین مرحله انجام شده",
    ],
    "village_type": ["village_type", "village type", "هدف/ اقماری", "هدف/اقماری"],
}

SITE_TYPE_MAP = {
    "new site": SiteType.NEW_SITE,
    "add tech": SiteType.ADD_TECH,
    "add technology": SiteType.ADD_TECH,
    "repeater": SiteType.REPEATER,
}

# Fields whose change on an existing record raises a Change Request (§2.4).
TRACKED_WORKITEM_FIELDS = (
    "requested_technology_raw",
    "assignment_date",
    "on_air_date",
    "last_site_status",
)


def _norm(text: object) -> str:
    """Normalize a header/value for matching: strip, collapse whitespace/newlines."""
    return re.sub(r"\s+", " ", str(text).replace("‌", " ")).strip().lower()


def parse_site_type(value: object) -> SiteType | None:
    if value is None:
        return None
    return SITE_TYPE_MAP.get(_norm(value))


def parse_village_type(value: object) -> VillageType | None:
    if value is None:
        return None
    s = _norm(value)
    if "اقمار" in s or "satellite" in s:
        return VillageType.SATELLITE
    if "هدف" in s or "target" in s:
        return VillageType.TARGET
    return None


def parse_technologies(value: object) -> dict[str, bool]:
    """`2G3G4G` / `3G4G` -> {requires_2g, requires_3g, requires_4g}."""
    s = str(value or "").upper()
    return {
        "requires_2g": "2G" in s,
        "requires_3g": "3G" in s,
        "requires_4g": "4G" in s,
    }


def parse_date(value: object) -> date | None:
    if value in (None, "", "-"):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%d/%m/%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(str(value).strip(), fmt).date()
        except ValueError:
            continue
    return None


def _clean(value: object) -> str | None:
    if value in (None, "", "-"):
        return None
    return str(value).strip()


@dataclass
class RowRecord:
    site_id: str | None
    temp_site_id: str | None
    province: str | None
    city: str | None
    village_id: str | None
    village_name: str | None
    site_type: SiteType | None
    requested_technology_raw: str | None
    site_assignment_date: date | None
    on_air_date: date | None
    last_site_status: str | None
    village_type: VillageType | None
    row_number: int

    @property
    def site_code(self) -> str | None:
        if self.site_id and self.site_id != NO_SITE_ID:
            return self.site_id
        return self.temp_site_id

    @property
    def is_temporary(self) -> bool:
        return not (self.site_id and self.site_id != NO_SITE_ID)


@dataclass
class ImportResult:
    rows_total: int = 0
    sites_created: int = 0
    work_items_created: int = 0
    villages_created: int = 0
    unchanged: int = 0
    change_requests: list[dict] = field(default_factory=list)
    errors: list[dict] = field(default_factory=list)

    def as_summary(self) -> dict:
        return {
            "rows_total": self.rows_total,
            "sites_created": self.sites_created,
            "work_items_created": self.work_items_created,
            "villages_created": self.villages_created,
            "unchanged": self.unchanged,
            "change_requests": len(self.change_requests),
            "errors": len(self.errors),
        }


def detect_header_row(ws, forced: int | None = None, scan_limit: int = 15) -> int:
    """Return the 1-based row index that contains the column headers.

    Detection looks for the row that contains a ``site_id`` alias.
    """
    if forced:
        return forced
    site_aliases = {a for a in COLUMN_ALIASES["site_id"]}
    for r in range(1, min(scan_limit, ws.max_row) + 1):
        values = {_norm(ws.cell(row=r, column=c).value) for c in range(1, ws.max_column + 1)}
        if values & site_aliases:
            return r
    raise ValueError("Could not locate CPM header row (no 'site_id' column found).")


def build_column_map(ws, header_row: int) -> dict[str, int]:
    """Map canonical field name -> 1-based column index (first match wins)."""
    headers = {
        c: _norm(ws.cell(row=header_row, column=c).value)
        for c in range(1, ws.max_column + 1)
    }
    col_map: dict[str, int] = {}
    for field_name, aliases in COLUMN_ALIASES.items():
        norm_aliases = [_norm(a) for a in aliases]
        for col, header in headers.items():
            if header and header in norm_aliases and field_name not in col_map:
                col_map[field_name] = col
                break
    return col_map


def read_rows(path: str | Path, header_row: int | None = None) -> list[RowRecord]:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    hr = detect_header_row(ws, header_row)
    cmap = build_column_map(ws, hr)

    required = {"site_id", "site_type", "village_id"}
    missing = required - cmap.keys()
    if missing:
        raise ValueError(f"CPM file is missing required columns: {sorted(missing)}")

    def cell(row: int, field_name: str):
        col = cmap.get(field_name)
        return ws.cell(row=row, column=col).value if col else None

    records: list[RowRecord] = []
    for r in range(hr + 1, ws.max_row + 1):
        raw = {f: cell(r, f) for f in COLUMN_ALIASES}
        if all(v in (None, "") for v in raw.values()):
            continue
        records.append(
            RowRecord(
                site_id=_clean(raw["site_id"]),
                temp_site_id=_clean(raw["temp_site_id"]),
                province=_clean(raw["province"]),
                city=_clean(raw["city"]),
                village_id=_clean(raw["village_id"]),
                village_name=_clean(raw["village_name"]),
                site_type=parse_site_type(raw["site_type"]),
                requested_technology_raw=_clean(raw["requested_technology"]),
                site_assignment_date=parse_date(raw["site_assignment_date"]),
                on_air_date=parse_date(raw["on_air_date"]),
                last_site_status=_clean(raw["last_site_status"]),
                village_type=parse_village_type(raw["village_type"]),
                row_number=r,
            )
        )
    return records


class CpmImportService:
    """Orchestrates a CPM import against the database (§3.21)."""

    def __init__(self, db: Session):
        self.db = db
        self._province_cache: dict[str, Province] = {}
        self._city_cache: dict[tuple[int, str], City] = {}

    # --- master data get-or-create ---
    def _get_province(self, name: str) -> Province:
        if name not in self._province_cache:
            prov = (
                self.db.query(Province).filter(Province.name == name).one_or_none()
            )
            if prov is None:
                prov = Province(name=name)
                self.db.add(prov)
                self.db.flush()
            self._province_cache[name] = prov
        return self._province_cache[name]

    def _get_city(self, province: Province, name: str | None) -> City | None:
        if not name:
            return None
        key = (province.id, name)
        if key not in self._city_cache:
            city = (
                self.db.query(City)
                .filter(City.province_id == province.id, City.name == name)
                .one_or_none()
            )
            if city is None:
                city = City(name=name, province_id=province.id)
                self.db.add(city)
                self.db.flush()
            self._city_cache[key] = city
        return self._city_cache[key]

    def _find_site(self, rec: RowRecord) -> Site | None:
        # Match by effective code first; then detect a temp→real promotion.
        site = (
            self.db.query(Site)
            .filter(Site.site_code == rec.site_code, Site.is_deleted.is_(False))
            .one_or_none()
        )
        if site:
            return site
        if not rec.is_temporary and rec.temp_site_id:
            return (
                self.db.query(Site)
                .filter(Site.temp_site_id == rec.temp_site_id, Site.is_deleted.is_(False))
                .one_or_none()
            )
        return None

    def import_file(
        self, path: str | Path, filename: str | None = None,
        header_row: int | None = None, imported_by: int | None = None,
    ) -> tuple[CpmImportBatch, ImportResult]:
        records = read_rows(path, header_row)
        result = ImportResult(rows_total=len(records))
        batch = CpmImportBatch(filename=filename, imported_by=imported_by)
        self.db.add(batch)
        self.db.flush()

        for rec in records:
            try:
                self._process_row(rec, batch, result)
            except Exception as exc:  # noqa: BLE001 — capture, keep importing
                result.errors.append({"row": rec.row_number, "error": str(exc)})

        batch.rows_total = result.rows_total
        batch.sites_created = result.sites_created
        batch.work_items_created = result.work_items_created
        batch.villages_created = result.villages_created
        batch.unchanged = result.unchanged
        batch.change_requests = len(result.change_requests)
        batch.errors = len(result.errors)
        self.db.commit()
        return batch, result

    def _process_row(self, rec: RowRecord, batch: CpmImportBatch, result: ImportResult) -> None:
        if not rec.site_code or rec.site_type is None or not rec.village_id:
            raise ValueError("missing site_code / site_type / village_id")
        if not rec.province:
            raise ValueError("missing province")

        province = self._get_province(rec.province)
        city = self._get_city(province, rec.city)

        # --- Site ---
        site = self._find_site(rec)
        if site is None:
            site = Site(
                site_code=rec.site_code,
                irancell_site_id=None if rec.is_temporary else rec.site_id,
                temp_site_id=rec.temp_site_id,
                is_temporary=rec.is_temporary,
                province_id=province.id,
                city_id=city.id if city else None,
            )
            self.db.add(site)
            self.db.flush()
            result.sites_created += 1
        elif site.is_temporary and not rec.is_temporary:
            # temp code promoted to a real Site ID -> Change Request (§2.4)
            self._raise_change(
                batch, result, "Site", site.site_code, "irancell_site_id",
                old=None, new=rec.site_id,
            )

        # --- Work Item (Site + Site Type) ---
        wi = (
            self.db.query(WorkItem)
            .filter(
                WorkItem.site_id == site.id,
                WorkItem.site_type == rec.site_type,
                WorkItem.is_deleted.is_(False),
            )
            .one_or_none()
        )
        techs = parse_technologies(rec.requested_technology_raw)
        if wi is None:
            wi = WorkItem(
                site_id=site.id,
                site_type=rec.site_type,
                requested_technology_raw=rec.requested_technology_raw,
                assignment_date=rec.site_assignment_date,
                on_air_date=rec.on_air_date,
                last_site_status=rec.last_site_status,
                **techs,
            )
            self.db.add(wi)
            self.db.flush()
            result.work_items_created += 1
        else:
            self._detect_workitem_changes(wi, rec, batch, result)

        # --- Village ---
        village = (
            self.db.query(Village)
            .filter(
                Village.work_item_id == wi.id,
                Village.village_code == rec.village_id,
                Village.is_deleted.is_(False),
            )
            .one_or_none()
        )
        if village is None:
            village = Village(
                work_item_id=wi.id,
                village_code=rec.village_id,
                name=rec.village_name,
                village_type=rec.village_type,
            )
            self.db.add(village)
            self.db.flush()
            self._materialize_approvals(village, techs)
            result.villages_created += 1
        else:
            result.unchanged += 1

    def _detect_workitem_changes(
        self, wi: WorkItem, rec: RowRecord, batch: CpmImportBatch, result: ImportResult
    ) -> None:
        incoming = {
            "requested_technology_raw": rec.requested_technology_raw,
            "assignment_date": rec.site_assignment_date,
            "on_air_date": rec.on_air_date,
            "last_site_status": rec.last_site_status,
        }
        changed = False
        for f in TRACKED_WORKITEM_FIELDS:
            old = getattr(wi, f)
            new = incoming[f]
            if str(old) != str(new) and new is not None:
                self._raise_change(
                    batch, result, "WorkItem",
                    f"{wi.site.site_code}|{wi.site_type.value}", f,
                    old=old, new=new,
                )
                changed = True
        if not changed:
            result.unchanged += 1

    def _raise_change(self, batch, result, entity_type, entity_key, field_name, old, new):
        cr = ChangeRequest(
            batch_id=batch.id,
            entity_type=entity_type,
            entity_key=entity_key,
            field=field_name,
            old_value=None if old is None else str(old),
            new_value=None if new is None else str(new),
        )
        self.db.add(cr)
        result.change_requests.append(
            {"entity": entity_type, "key": entity_key, "field": field_name,
             "old": str(old), "new": str(new)}
        )

    def _materialize_approvals(self, village: Village, techs: dict[str, bool]) -> None:
        """Pre-create PENDING approvals per requested technology, ICT & CRA (§3.11)."""
        wanted = []
        if techs["requires_2g"]:
            wanted.append(Technology.G2)
        if techs["requires_3g"]:
            wanted.append(Technology.G3)
        if techs["requires_4g"]:
            wanted.append(Technology.G4)
        for phase in ("ICT", "CRA"):
            for tech in wanted:
                self.db.add(
                    VillageTechApproval(
                        village_id=village.id,
                        phase=phase,
                        technology=tech,
                        status=ApprovalStatus.PENDING,
                    )
                )
